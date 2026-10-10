"""Extract tariff cells for the response panel from the user-provided workbook.

Run after replacing HS_AI/TL/BIEU THUE XNK 2026.xlsx. The generated JSON is
source data, not a declaration that every rate is currently in force.
"""

import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'HS_AI' / 'TL' / 'BIEU THUE XNK 2026.xlsx'
TARGET = ROOT / 'backend' / 'src' / 'hs-codes' / 'tariff-rates.json'


def cell(value):
    return str(value).strip() if value is not None and str(value).strip() else None


sheet = load_workbook(SOURCE, read_only=True, data_only=True)['BT2026']
rates = {}
for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
    code = cell(row[5]) if len(row) > 5 else None
    if not code or not code.isdigit() or len(code) != 8:
        continue
    # Sheet columns: K normal, N MFN, Q VAT, T ACFTA, W ATIGA, AF AKFTA.
    values = [cell(row[index]) if len(row) > index else None for index in (10, 13, 16, 19, 22, 31)]
    if any(value is not None for value in values):
        rates[code] = [row_number, *values]

TARGET.write_text(json.dumps(rates, ensure_ascii=False, separators=(',', ':')) + '\n', encoding='utf-8')
print(f'Wrote {len(rates)} tariff rows to {TARGET}')
