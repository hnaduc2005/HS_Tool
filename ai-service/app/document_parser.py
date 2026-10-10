import csv
import re
import unicodedata
from pathlib import Path
from typing import Iterator

def clean(text: str) -> str:
    text = unicodedata.normalize('NFC', text)
    text = text.replace('\u00ad', '').replace('\u200b', '').replace('\ufeff', '').replace('\u00a0', ' ')
    text = ''.join(c for c in text if c in '\n\t' or unicodedata.category(c)[0] != 'C')
    return re.sub(r'[ \t]+', ' ', text).strip()

def extract(path: Path) -> Iterator[tuple[str, int | None, dict]]:
    suffix = path.suffix.lower()
    if suffix == '.pdf':
        import pymupdf
        with pymupdf.open(path) as pdf:
            for page_no, page in enumerate(pdf, 1):
                content = clean(page.get_text())
                if content: yield content, page_no, {'file': path.name, 'page': page_no}
    elif suffix == '.docx':
        from docx import Document
        doc = Document(path)
        content = clean('\n'.join(p.text for p in doc.paragraphs))
        if content: yield content, None, {'file': path.name}
        for index, table in enumerate(doc.tables, 1):
            content = clean('\n'.join(' | '.join(cell.text for cell in row.cells) for row in table.rows))
            if content: yield content, None, {'file': path.name, 'table': index}
    elif suffix == '.xlsx':
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in workbook:
                for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    content = clean(' | '.join(str(value) for value in row if value is not None))
                    if content: yield content, None, {'file': path.name, 'sheet': sheet.title, 'row': row_number}
        finally: workbook.close()
    elif suffix == '.txt':
        content = clean(path.read_text(encoding='utf-8-sig'))
        if content: yield content, None, {'file': path.name}
    else:
        raise ValueError('Unsupported document type')

def legal_extract(path: Path) -> Iterator[tuple[str, int | None, dict]]:
    """Keep section, chapter and heading context across pages of an HS reference."""
    context: dict = {}
    marker = re.compile(r'(?mi)^[ \t]*(?:Chương[ \t]+(?P<chapter>\d{1,2})\b|Phần[ \t]+(?P<section>[IVX]+)\b|(?P<heading>\d{2}\.\d{2})\.?[ \t]*[-–])')
    for content, page, metadata in extract(path):
        boundaries = list(marker.finditer(content))
        start = 0
        for index, match in enumerate(boundaries):
            if match.start() > start:
                yield content[start:match.start()], page, {**metadata, **context}
            if match.group('section'):
                context = {'hs_section': match.group('section').upper(), 'legal_kind': 'section_note'}
            elif match.group('chapter'):
                context = {**context, 'hs_chapter': match.group('chapter').zfill(2), 'hs_heading': None, 'legal_kind': 'chapter_note'}
            else:
                heading = match.group('heading').replace('.', '')
                context = {**context, 'hs_chapter': heading[:2], 'hs_heading': heading, 'legal_kind': 'heading'}
            end = boundaries[index + 1].start() if index + 1 < len(boundaries) else len(content)
            yield content[match.start():end], page, {**metadata, **context}
            start = end
        if not boundaries:
            yield content, page, {**metadata, **context}


def chunks(path: Path, max_chars: int = 1200, overlap: int = 150) -> Iterator[tuple[str, int | None, dict]]:
    source = legal_extract(path) if path.suffix.lower() == '.pdf' else extract(path)
    for content, page, metadata in source:
        start = 0
        while start < len(content):
            end = min(start + max_chars, len(content))
            if end < len(content):
                # Prefer a paragraph/sentence boundary to keep definitions together.
                lower = start + max_chars // 2
                boundaries = [match.end() for match in re.finditer(r'\n\s*\n|[.;:]\s+', content[lower:end])]
                boundary = lower + boundaries[-1] if boundaries else content.rfind(' ', lower, end)
                if boundary > start: end = boundary
            part = content[start:end].strip()
            if part: yield part, page, metadata
            if end == len(content): break
            start = max(start + 1, end - overlap)

def parse_hs(path: Path) -> list[dict]:
    rows = []
    if path.suffix.lower() == '.xlsx':
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook['BT2026'] if 'BT2026' in workbook.sheetnames else workbook.active
        iterable = ((row[5], row[6]) for row in sheet.iter_rows(values_only=True) if len(row) > 6)
    elif path.suffix.lower() == '.csv':
        file = path.open('r', encoding='utf-8-sig', newline='')
        reader = csv.DictReader(file)
        iterable = ((row.get('code'), row.get('description')) for row in reader)
    else:
        raise ValueError('Unsupported HS import format')
    try:
        seen = set()
        uncoded_labels: dict[int, str] = {}
        for code, description in iterable:
            value = str(code or '').replace('.', '').strip()
            if path.suffix.lower() == '.xlsx' and description:
                label = clean(str(description))
                match = re.match(r'^((?:-\s*)+)(.*)', label)
                depth = match.group(1).count('-') if match else 0
                uncoded_labels = {level: text for level, text in uncoded_labels.items() if level < depth}
                if not value and match:
                    uncoded_labels[depth] = match.group(2).strip()
                    continue
                description = ' → '.join([*uncoded_labels.values(), match.group(2).strip() if match else label])
            if not re.fullmatch(r'\d{4}(?:\d{2})?(?:\d{2})?', value) or not description or value in seen: continue
            seen.add(value)
            rows.append({'code': value, 'description': clean(str(description)), 'level': len(value), 'parent_code': value[:len(value)-2] if len(value)>4 else None})
    finally:
        if path.suffix.lower() == '.xlsx': workbook.close()
        else: file.close()
    return rows
