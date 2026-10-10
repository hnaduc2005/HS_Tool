from pathlib import Path
from app.document_parser import clean, chunks, parse_hs

FIXTURES = Path(__file__).parent / 'fixtures'

def test_clean_removes_controls():
    assert clean('  Áo\x00   thun\nnam ') == 'Áo thun\nnam'

def test_txt_chunks_preserve_traceability():
    parts = list(chunks(FIXTURES / 'sample.txt', max_chars=100, overlap=10))
    assert len(parts) > 1
    assert all(item[2]['file'] == 'sample.txt' for item in parts)

def test_parse_hs_csv_filters_invalid_rows():
    rows = parse_hs(FIXTURES / 'codes.csv')
    assert rows == [{'code': '61091010', 'description': 'Áo phông nam', 'level': 8, 'parent_code': '610910'}]


def test_xlsx_preserves_uncoded_parent_labels_and_resets_sibling_context(monkeypatch):
    from io import BytesIO
    from openpyxl import Workbook, load_workbook
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'BT2026'
    for code, description in [
        ('160249', '- - Loại khác, kể cả các hỗn hợp:'),
        (None, '- - - Thịt nguội:'),
        ('16024911', '- - - - Đóng bao bì kín khí để bán lẻ'),
        (None, '- - - Loại khác:'),
        ('16024991', '- - - - Đóng bao bì kín khí để bán lẻ'),
        ('160529', '- - Loại khác:'),
        ('16052990', '- - - Loại khác'),
    ]:
        sheet.append([None] * 5 + [code, description])
    data = BytesIO()
    workbook.save(data)
    data.seek(0)
    monkeypatch.setattr('openpyxl.load_workbook', lambda path, **options: load_workbook(data, **options))
    rows = {row['code']: row for row in parse_hs(Path('catalog.xlsx'))}
    assert rows['16024911']['description'] == 'Thịt nguội: → Đóng bao bì kín khí để bán lẻ'
    assert rows['16024991']['description'] == 'Loại khác: → Đóng bao bì kín khí để bán lẻ'
    assert rows['16052990']['description'] == 'Loại khác'
